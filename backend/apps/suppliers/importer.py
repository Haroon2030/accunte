"""
استيراد الموردين من Excel: تحليل ← تحقق ← مرحلة تجهيز (dry-run) ← كتابة نهائية.
"""
import io
import re

from django.contrib.auth import get_user_model
from django.db import transaction
from django.utils import timezone
from openpyxl import Workbook, load_workbook

from apps.audit.models import AuditLog
from apps.audit.services import log_action
from apps.branches.models import Branch

from .models import Supplier, SupplierContact, SupplierImportJob

MAX_ROWS = 5000
CHUNK = 500

# اسم الحقل الداخلي ← العناوين المقبولة في الملف (بعد التطبيع)
COLUMNS = {
    'code': ['رقم المورد', 'الكود', 'code', 'supplier code', 'supplier_code'],
    'financial_system_id': ['رقم المورد في النظام المالي', 'رقم النظام المالي', 'financial_system_id', 'financial id'],
    'name': ['اسم المورد', 'الاسم', 'name', 'supplier name', 'الاسم التجاري'],
    'tax_number': ['الرقم الضريبي', 'tax_number', 'vat', 'vat number', 'tax id'],
    'phone': ['الهاتف', 'الجوال', 'رقم الهاتف', 'ارقام الهاتف', 'أرقام الهاتف', 'phone', 'mobile'],
    'email': ['البريد الإلكتروني', 'البريد', 'email'],
    'contact_person': ['الشخص المسؤول', 'جهة الاتصال', 'المسؤول', 'contact', 'contact person'],
    'branches': ['الفروع', 'الفروع المرتبطة', 'branches', 'branch'],
    'accountant': ['المحاسب', 'المحاسب المسؤول', 'accountant'],
    'notes': ['ملاحظات', 'notes'],
    'bank_name': ['اسم البنك', 'البنك', 'bank', 'bank name'],
    'iban': ['الآيبان', 'الايبان', 'iban'],
}
TEMPLATE_HEADERS = ['رقم المورد', 'رقم المورد في النظام المالي', 'اسم المورد', 'الرقم الضريبي',
                    'الهاتف', 'البريد الإلكتروني', 'الشخص المسؤول', 'الفروع', 'المحاسب', 'ملاحظات']

ALIAS_TO_FIELD = {}
for _field, _aliases in COLUMNS.items():
    for _a in _aliases:
        ALIAS_TO_FIELD[_a.strip().lower()] = _field

EMAIL_RE = re.compile(r'^[^@\s]+@[^@\s]+\.[^@\s]+$')


def normalize_phone(raw):
    """يحوّل الأرقام السعودية إلى صيغة دولية 9665XXXXXXXX. يُرجع None إن كان غير صالح."""
    digits = re.sub(r'\D', '', str(raw or ''))
    if not digits:
        return None
    if digits.startswith('00'):
        digits = digits[2:]
    if digits.startswith('05') and len(digits) == 10:
        digits = '966' + digits[1:]
    elif digits.startswith('5') and len(digits) == 9:
        digits = '966' + digits
    if not 8 <= len(digits) <= 15:
        return None
    return digits


def split_multi(value):
    return [p.strip() for p in re.split(r'[,;،؛/\n]+', str(value or '')) if p.strip()]


def _cell(value):
    if value is None:
        return ''
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


def build_template():
    wb = Workbook()
    ws = wb.active
    ws.title = 'الموردين'
    ws.sheet_view.rightToLeft = True
    ws.append(TEMPLATE_HEADERS)
    ws.append(['S-0001', 'FIN-1001', 'شركة المثال للتجارة', '300000000000003',
               '0501234567, 0559876543', 'info@example.com', 'أحمد محمد', 'RYD, JED', 'accountant1', ''])
    for col in 'ABCDEFGHIJ':
        ws.column_dimensions[col].width = 24
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


class ImportFileError(Exception):
    pass


def read_rows(file_obj):
    """يقرأ ملف xlsx ويُرجع [(رقم_الصف, {field: value})]."""
    try:
        wb = load_workbook(file_obj, read_only=True, data_only=True)
    except Exception as exc:
        raise ImportFileError('تعذر قراءة الملف. تأكد أنه بصيغة Excel (.xlsx)') from exc
    ws = wb.active
    rows_iter = ws.iter_rows(values_only=True)
    try:
        header = next(rows_iter)
    except StopIteration:
        raise ImportFileError('الملف فارغ')
    mapping = {}
    for idx, title in enumerate(header):
        field = ALIAS_TO_FIELD.get(_cell(title).lower())
        if field and field not in mapping.values():
            mapping[idx] = field
    if 'name' not in mapping.values() or ('code' not in mapping.values() and 'financial_system_id' not in mapping.values()):
        raise ImportFileError('الأعمدة المطلوبة غير موجودة: "اسم المورد" و"رقم المورد" (أو رقم النظام المالي). حمّل القالب للتأكد من العناوين.')
    out = []
    for n, row in enumerate(rows_iter, start=2):
        data = {f: _cell(row[i]) if i < len(row) else '' for i, f in mapping.items()}
        if not any(data.values()):
            continue
        out.append((n, data))
        if len(out) > MAX_ROWS:
            raise ImportFileError(f'عدد الصفوف يتجاوز الحد الأقصى ({MAX_ROWS})')
    return out


def stage(file_obj, filename, user):
    """يحلل الملف ويحفظ نتيجة التحقق في SupplierImportJob دون لمس جدول الموردين."""
    raw = read_rows(file_obj)

    db_codes = set(Supplier.objects.values_list('code', flat=True))
    db_fin = set(Supplier.objects.exclude(financial_system_id__isnull=True)
                 .exclude(financial_system_id='').values_list('financial_system_id', flat=True))
    db_tax = set(Supplier.objects.exclude(tax_number='').values_list('tax_number', flat=True))
    branch_map = {}
    for b in Branch.objects.all():
        branch_map[b.code.lower()] = b.pk
        branch_map[b.name.strip().lower()] = b.pk
    User = get_user_model()
    user_map = {u.get_username().lower(): u.pk for u in User.objects.filter(is_active=True)}

    seen_codes, seen_fin, seen_tax = set(), set(), set()
    rows = []
    counts = {'valid': 0, 'duplicate': 0, 'error': 0}

    for n, d in raw:
        errors, dups = [], []
        code = d.get('code') or d.get('financial_system_id', '')
        fin = d.get('financial_system_id', '')
        tax = re.sub(r'\s', '', d.get('tax_number', ''))
        name = d.get('name', '')

        if not name:
            errors.append('اسم المورد مطلوب')
        if not code:
            errors.append('رقم المورد مطلوب')
        if len(code) > 50 or len(fin) > 50:
            errors.append('رقم المورد أطول من 50 حرفًا')
        if tax and not re.fullmatch(r'\d{15}', tax):
            errors.append('الرقم الضريبي يجب أن يكون 15 رقمًا')
        email = d.get('email', '')
        if email and not EMAIL_RE.match(email):
            errors.append('بريد إلكتروني غير صالح')

        phones, bad_phones = [], []
        for raw_phone in split_multi(d.get('phone', '')):
            p = normalize_phone(raw_phone)
            (phones if p else bad_phones).append(p or raw_phone)
        if bad_phones:
            errors.append('رقم هاتف غير صالح: ' + '، '.join(bad_phones))
        phones = list(dict.fromkeys(phones))

        branch_ids, bad_branches = [], []
        for token in split_multi(d.get('branches', '')):
            bid = branch_map.get(token.lower())
            (branch_ids if bid else bad_branches).append(bid or token)
        if bad_branches:
            errors.append('فرع غير معروف: ' + '، '.join(bad_branches))

        accountant_id = None
        if d.get('accountant'):
            accountant_id = user_map.get(d['accountant'].lower())
            if accountant_id is None:
                errors.append('المحاسب غير موجود: ' + d['accountant'])

        if code:
            if code in db_codes:
                dups.append('رقم المورد موجود مسبقًا في النظام')
            elif code in seen_codes:
                dups.append('رقم المورد مكرر داخل الملف')
        if fin:
            if fin in db_fin:
                dups.append('رقم النظام المالي موجود مسبقًا')
            elif fin in seen_fin:
                dups.append('رقم النظام المالي مكرر داخل الملف')
        if tax:
            if tax in db_tax:
                dups.append('الرقم الضريبي موجود مسبقًا')
            elif tax in seen_tax:
                dups.append('الرقم الضريبي مكرر داخل الملف')

        status = 'error' if errors else 'duplicate' if dups else 'valid'
        counts[status] += 1
        if status == 'valid':   # الصفوف المقبولة فقط تحجز المعرفات
            seen_codes.add(code)
            if fin:
                seen_fin.add(fin)
            if tax:
                seen_tax.add(tax)

        rows.append({
            'row': n, 'status': status, 'messages': errors + dups,
            'name': name, 'code': code, 'financial_system_id': fin, 'tax_number': tax,
            'phones': phones, 'email': email, 'contact_person': d.get('contact_person', ''),
            'branch_ids': branch_ids, 'accountant_id': accountant_id,
            'notes': d.get('notes', ''), 'bank_name': d.get('bank_name', ''), 'iban': d.get('iban', ''),
        })

    return SupplierImportJob.objects.create(
        uploaded_by=user, filename=filename[:255], rows=rows, total_rows=len(rows),
        valid_rows=counts['valid'], duplicate_rows=counts['duplicate'], error_rows=counts['error'],
    )


@transaction.atomic
def commit(job, user=None):
    """يكتب الصفوف الصالحة فقط. يعيد فحص التكرار مقابل قاعدة البيانات لحظة التأكيد."""
    job = SupplierImportJob.objects.select_for_update().get(pk=job.pk)
    if job.status != SupplierImportJob.Status.STAGED:
        raise ValueError('هذه العملية غير قابلة للتأكيد')

    valid = [r for r in job.rows if r['status'] == 'valid']
    codes = [r['code'] for r in valid]
    existing_codes = set(Supplier.objects.filter(code__in=codes).values_list('code', flat=True))
    existing_fin = set(Supplier.objects.filter(
        financial_system_id__in=[r['financial_system_id'] for r in valid if r['financial_system_id']]
    ).values_list('financial_system_id', flat=True))
    existing_tax = set(Supplier.objects.filter(
        tax_number__in=[r['tax_number'] for r in valid if r['tax_number']]
    ).values_list('tax_number', flat=True))

    to_create, skipped = [], 0
    for r in valid:
        if r['code'] in existing_codes or (r['financial_system_id'] and r['financial_system_id'] in existing_fin) \
                or (r['tax_number'] and r['tax_number'] in existing_tax):
            skipped += 1   # سُجّل بعد المعاينة، تجاوزناه
            continue
        to_create.append(r)

    for i in range(0, len(to_create), CHUNK):
        part = to_create[i:i + CHUNK]
        Supplier.objects.bulk_create([
            Supplier(
                name=r['name'], code=r['code'], financial_system_id=r['financial_system_id'] or None,
                tax_number=r['tax_number'], tax_key=r['tax_number'] or None,
                phone=(r['phones'][0] if r['phones'] else ''),
                email=r['email'], contact_person=r['contact_person'], notes=r['notes'],
                bank_name=r['bank_name'], iban=r['iban'], assigned_accountant_id=r['accountant_id'],
            ) for r in part
        ])

    created_ids = dict(Supplier.objects.filter(code__in=[r['code'] for r in to_create]).values_list('code', 'pk'))
    contacts, links = [], []
    Through = Supplier.branches.through
    for r in to_create:
        sid = created_ids[r['code']]
        for idx, phone in enumerate(r['phones']):
            contacts.append(SupplierContact(supplier_id=sid, name=r['contact_person'], phone=phone,
                                            email=r['email'], is_primary=(idx == 0)))
        links.extend(Through(supplier_id=sid, branch_id=b) for b in dict.fromkeys(r['branch_ids']))
    SupplierContact.objects.bulk_create(contacts, batch_size=CHUNK, ignore_conflicts=True)
    Through.objects.bulk_create(links, batch_size=CHUNK, ignore_conflicts=True)

    job.status = SupplierImportJob.Status.COMMITTED
    job.created_count = len(to_create)
    job.committed_at = timezone.now()
    job.save(update_fields=['status', 'created_count', 'committed_at'])

    log_action(
        AuditLog.Action.IMPORT, entity=job, entity_repr=f'استيراد موردين: {job.filename}',
        new={'created': len(to_create), 'skipped_after_preview': skipped,
             'duplicates': job.duplicate_rows, 'errors': job.error_rows, 'total': job.total_rows},
        actor=user,
    )
    return job
