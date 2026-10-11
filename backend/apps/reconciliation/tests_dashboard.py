import datetime
import io
from decimal import Decimal

from django.contrib.auth import get_user_model
from django.test import TestCase
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from apps.audit.models import AuditLog
from apps.branches.models import Branch
from apps.core.models import Role

from . import kpis, services
from .models import ReconciliationCase as Case
from .models import ReconciliationDifference, ReconciliationResult, StatusHistory
from .tests import Base, make_user

S = Case.Status
DAY = datetime.timedelta(days=1)


class DashboardBase(Base):
    def setUp(self):
        super().setUp()
        self.today = timezone.localdate()
        self.b1 = Branch.objects.create(name='الرياض', code='RYD')
        self.b2 = Branch.objects.create(name='جدة', code='JED')

    def hist(self, case, to, days_ago):
        h = StatusHistory.objects.create(case=case, from_status='', to_status=to, changed_by=self.cfo)
        StatusHistory.objects.filter(pk=h.pk).update(created_at=timezone.now() - datetime.timedelta(days=days_ago))
        return h

    def age(self, case, days):
        Case.objects.filter(pk=case.pk).update(created_at=timezone.now() - datetime.timedelta(days=days))


class KpiDefinitionTests(DashboardBase):
    def test_counts_by_definition(self):
        self.new_case(status=S.NEW)
        self.new_case(status=S.READY_FOR_DISTRIBUTION)
        self.new_case(status=S.RECONCILING, assigned_accountant=self.acc)
        self.new_case(status=S.WAITING_VENDOR, assigned_accountant=self.acc)
        self.new_case(status=S.COMPLETED)
        self.new_case(status=S.CANCELLED)
        c = kpis.build(Case.objects.all())['counts']
        self.assertEqual((c['all'], c['new'], c['in_progress'], c['waiting_vendor'], c['completed']), (6, 2, 2, 1, 1))

    def test_overdue_excludes_terminal_and_future(self):
        self.new_case(status=S.RECONCILING, due_date=self.today - DAY, assigned_accountant=self.acc)
        self.new_case(status=S.COMPLETED, due_date=self.today - DAY)
        self.new_case(status=S.RECONCILING, due_date=self.today + DAY, assigned_accountant=self.acc)
        self.new_case(status=S.RECONCILING, due_date=None, assigned_accountant=self.acc)
        self.assertEqual(kpis.build(Case.objects.all())['counts']['overdue'], 1)

    def test_stale_uses_last_activity_not_creation(self):
        old_active = self.new_case(status=S.RECONCILING, assigned_accountant=self.acc)
        self.age(old_active, 60)
        self.hist(old_active, 'reconciling', 1)                 # قديمة لكن نشاطها أمس: ليست راكدة
        silent = self.new_case(status=S.RECONCILING, assigned_accountant=self.acc)
        self.hist(silent, 'reconciling', 20)
        self.hist(silent, 'reconciling', 15)                    # آخر نشاط قبل 15 يومًا: راكدة، وتُعدّ مرة واحدة
        no_history = self.new_case(status=S.ASSIGNED, assigned_accountant=self.acc)
        self.age(no_history, 30)                                # بلا سجل: يُقاس من الإنشاء
        done = self.new_case(status=S.COMPLETED)
        self.age(done, 90)                                      # المنتهية لا تُعدّ راكدة
        fresh = self.new_case(status=S.RECONCILING, assigned_accountant=self.acc)
        stale_ids = set(kpis.kpi_queryset('stale', Case.objects.all()).values_list('pk', flat=True))
        self.assertEqual(stale_ids, {silent.pk, no_history.pk})
        self.assertEqual(kpis.build(Case.objects.all())['counts']['stale'], 2)
        self.assertNotIn(fresh.pk, stale_ids)

    def test_avg_resolution_time_from_first_completion(self):
        a = self.new_case(status=S.COMPLETED)
        self.age(a, 10)
        self.hist(a, 'completed', 6)          # 4 أيام
        b = self.new_case(status=S.CLOSED)
        self.age(b, 20)
        self.hist(b, 'completed', 8)          # 12 يومًا
        self.hist(b, 'closed', 0)             # الإغلاق لا يغيّر زمن الحل
        self.new_case(status=S.RECONCILING, assigned_accountant=self.acc)   # غير منتهية: خارج المتوسط
        out = kpis.build(Case.objects.all())
        self.assertEqual(out['resolved_n'], 2)
        self.assertAlmostEqual(out['avg_resolution']['days'], 8.0, places=1)

    def test_avg_resolution_none_without_data(self):
        self.assertIsNone(kpis.build(Case.objects.all())['avg_resolution'])


class DedupAndScopeTests(DashboardBase):
    def test_case_with_many_differences_counts_once(self):
        c = self.new_case(status=S.DISCREPANCY_FOUND, assigned_accountant=self.acc)
        for i in range(5):
            ReconciliationDifference.objects.create(case=c, supplier_amount=100 + i, company_amount=50)
        clean = self.new_case(status=S.RECONCILING, assigned_accountant=self.acc)
        out = kpis.build(Case.objects.all())
        self.assertEqual(out['counts']['discrepancy'], 1)
        self.assertEqual(out['discrepancy']['cases'], 1)
        self.assertEqual(out['discrepancy']['worked'], 2)
        self.assertEqual(out['discrepancy']['rate'], 50.0)
        self.assertEqual(out['discrepancy']['unresolved_items'], 5)
        self.assertEqual(out['discrepancy']['unresolved_amount'], Decimal('260'))   # 50+51+52+53+54
        self.assertEqual(kpis.kpi_queryset('discrepancy', Case.objects.all()).count(), 1)
        self.assertNotIn(clean.pk, kpis.kpi_queryset('discrepancy', Case.objects.all()).values_list('pk', flat=True))

    def test_resolved_differences_leave_unresolved_total(self):
        c = self.new_case(status=S.RECONCILING, assigned_accountant=self.acc)
        ReconciliationDifference.objects.create(case=c, supplier_amount=100, company_amount=0, is_resolved=True)
        ReconciliationDifference.objects.create(case=c, supplier_amount=30, company_amount=10)
        d = kpis.build(Case.objects.all())['discrepancy']
        self.assertEqual((d['unresolved_items'], d['unresolved_amount']), (1, Decimal('20')))

    def test_multi_branch_case_counts_once_in_branch_scope(self):
        c = self.new_case(status=S.RECONCILING, assigned_accountant=self.acc)
        c.branches.set([self.b1, self.b2])
        only_jed = self.new_case(status=S.RECONCILING, assigned_accountant=self.acc)
        only_jed.branches.set([self.b2])
        self.assertEqual(kpis.build(Case.objects.all(), branch_id=self.b1.pk)['counts']['all'], 1)
        self.assertEqual(kpis.build(Case.objects.all(), branch_id=self.b2.pk)['counts']['all'], 2)
        self.assertEqual(kpis.build(Case.objects.all())['counts']['all'], 2)

    def test_period_scope(self):
        old = self.new_case(status=S.NEW)
        self.age(old, 200)
        self.new_case(status=S.NEW)
        self.assertEqual(kpis.build(Case.objects.all(), '30')['counts']['all'], 1)
        self.assertEqual(kpis.build(Case.objects.all(), '365')['counts']['all'], 2)
        self.assertEqual(kpis.build(Case.objects.all(), 'all')['counts']['all'], 2)

    def test_proposal_totals_not_inflated_by_decisions(self):
        from apps.proposals.models import PaymentApproval, PaymentProposal
        case = self.new_case(status=S.COMPLETED)
        p = PaymentProposal.objects.create(
            supplier=self.supplier, branch=self.b1, reconciliation_case=case, total_dues=1000, proposed_amount=400,
            justification='x', created_by=self.acc, status='pending_review')
        for _ in range(3):
            PaymentApproval.objects.create(proposal=p, decision='needs_info', decided_by=self.cfo, amount_at_decision=400)
        pr = kpis.build(Case.objects.all())['proposals']
        self.assertEqual((pr['grand_count'], pr['grand_total']), (1, Decimal('400.00')))
        self.assertEqual((pr['pending_count'], pr['pending_total']), (1, Decimal('400.00')))




class DashboardAccessTests(DashboardBase):
    def test_roles(self):
        url = reverse('reconciliation:executive')
        for user, code in ((self.cfo, 200), (make_user('adm', Role.RoleType.ADMIN), 200),
                           (self.acc, 302), (make_user('sup_x', Role.RoleType.CS_SUPERVISOR), 302),
                           (make_user('vendor', Role.RoleType.SUPPLIER), 302)):
            self.client.force_login(user)
            self.assertEqual(self.client.get(url).status_code, code, user.username)
        self.client.logout()
        self.assertEqual(self.client.get(url).status_code, 302)

    def test_dashboard_renders_with_data(self):
        self.new_case(status=S.RECONCILING, due_date=self.today - DAY, assigned_accountant=self.acc)
        self.client.force_login(self.cfo)
        html = self.client.get(reverse('reconciliation:executive')).content.decode()
        self.assertIn('لوحة المدير المالي', html)
        self.assertIn('متأخرة', html)

    def test_export_endpoints_require_executive_role_for_dashboard(self):
        self.client.force_login(self.acc)
        for name in ('executive_excel', 'executive_print'):
            self.assertEqual(self.client.get(reverse(f'reconciliation:{name}')).status_code, 302)


class ExportTests(DashboardBase):
    def test_dashboard_excel_has_all_sheets_and_is_audited(self):
        self.new_case(status=S.RECONCILING, assigned_accountant=self.acc, due_date=self.today - DAY)
        self.client.force_login(self.cfo)
        resp = self.client.get(reverse('reconciliation:executive_excel'), {'period': '90'})
        self.assertEqual(resp.status_code, 200)
        self.assertIn('spreadsheetml', resp['Content-Type'])
        wb = load_workbook(io.BytesIO(resp.content))
        self.assertEqual(wb.sheetnames, ['المؤشرات', 'الحالات حسب الحالة', 'أحمال المحاسبين', 'مقترحات الدفع', 'متأخرة', 'راكدة'])
        cells = {r[0]: r[1] for r in wb['المؤشرات'].iter_rows(min_row=1, values_only=True) if r[0]}
        self.assertEqual(cells['متأخرة'], 1)
        self.assertTrue(AuditLog.objects.filter(action=AuditLog.Action.EXPORT, entity_repr__contains='Excel').exists())






class RequestStageRowsTests(TestCase):
    """بطاقة «توزيع الطلبات حسب المرحلة» تعدّ طلبات الموردين لا حالات المطابقة."""

    def test_stage_rows_count_requests_by_intake_stage(self):
        from apps.core.models import Role, UserProfile
        from apps.servicedesk.models import ServiceRequest, ServiceType
        from apps.suppliers.models import Supplier
        creator = get_user_model().objects.create_user('rr_creator', password='x')
        acc = get_user_model().objects.create_user('rr_acc', password='x')
        role, _ = Role.objects.get_or_create(role_type=Role.RoleType.ACCOUNTANT, defaults={'name': 'محاسب'})
        UserProfile.objects.update_or_create(user=acc, defaults={'role': role})
        sup = Supplier.objects.create(name='مورد', code='RR1')
        stype = ServiceType.objects.get(code='payment-status')
        mk = lambda **kw: ServiceRequest.objects.create(service_type=stype, supplier=sup, subject='x', created_by=creator, **kw)
        mk()                                                         # بانتظار الاستلام
        mk(assigned_to=creator)                                      # تم الاستلام
        mk(assigned_to=acc)                                          # أُرسلت للمطابق
        mk(status='in_progress', assigned_to=acc)
        mk(status='resolved', assigned_to=acc)
        rows, top = kpis.request_stage_rows('all')
        counts = {k: n for k, _, n in rows}
        self.assertEqual((counts['awaiting'], counts['received'], counts['sent'], counts['working'], counts['done']),
                         (1, 1, 1, 1, 1))
        self.assertEqual(top, 1)
