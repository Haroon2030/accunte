"""يولّد ملفات الاختبار الواقعية (Excel/CSV) المستخدمة في integration.test.js. التشغيل: python make_fixtures.py"""
import datetime as dt
import os

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

HERE = os.path.dirname(os.path.abspath(__file__))
D = dt.date


def ours():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'كشف حساب'
    ws.sheet_view.rightToLeft = True
    ws.merge_cells('A1:F1')
    ws['A1'] = 'شركة محمد رشيد الرشيد للتجارة'
    ws['A1'].font = Font(bold=True, size=14)
    ws['A2'] = 'كشف حساب مورد'
    ws['A3'] = 'اسم المورد: شركة الأفق للتوريدات'
    ws['A4'] = 'الفترة من 01/01/2026 إلى 31/01/2026'
    header = ['التاريخ', 'رقم المستند', 'البيان', 'مدين', 'دائن', 'الرصيد']
    for i, h in enumerate(header, 1):
        c = ws.cell(row=6, column=i, value=h)
        c.font = Font(bold=True, color='FFFFFF')
        c.fill = PatternFill('solid', fgColor='13314A')
        c.alignment = Alignment(horizontal='center')
    rows = [
        (None, None, 'رصيد افتتاحي', None, 5000, 5000),
        (D(2026, 1, 5), 'INV-1001', 'فاتورة مشتريات', None, 10000, 15000),
        (D(2026, 1, 12), 'PAY-77', 'سداد دفعة بنكية', 4000, None, 11000),
        (D(2026, 1, 20), 'INV-1002', 'فاتورة مشتريات', None, 2500.5, 13500.5),
        (D(2026, 1, 26), 'CN-5', 'إشعار دائن مرتجع', 700, None, 12800.5),
        (D(2026, 1, 28), 'PAY-80', 'سداد دفعة بنكية', 3000, None, 9800.5),
        (None, None, 'الإجمالي', 7700, 12500.5, 9800.5),
    ]
    for r, row in enumerate(rows, 7):
        for c, v in enumerate(row, 1):
            cell = ws.cell(row=r, column=c, value=v)
            if isinstance(v, dt.date):
                cell.number_format = 'dd/mm/yyyy'
            elif isinstance(v, (int, float)):
                cell.number_format = '#,##0.00'
    wb.save(os.path.join(HERE, 'fixture_ours.xlsx'))


def supplier():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Statement'
    ws['A1'] = 'كشف حساب العميل: شركة محمد رشيد الرشيد للتجارة'
    ws['A2'] = 'الرصيد السابق: 5,000.00'
    ws.merge_cells('D4:E4')
    ws['D4'] = 'المبلغ'
    for i, h in enumerate(['التاريخ', 'المرجع', 'الوصف', 'مدين', 'دائن', 'الرصيد التراكمي'], 1):
        ws.cell(row=5, column=i, value=h)
    # كشف بترتيب تنازلي (الأحدث أولًا): نبني الأرصدة من الأقدم ثم نعكس الصفوف
    chrono = [('05/01/2026', 'INV-1001', 'فاتورة مشتريات', 10000, None), ('06/01/2026', 'A1', 'فاتورة مشتريات', 1150, None),
              ('12/01/2026', 'PAY-77', 'دفعة مستلمة تحويل', None, 4000), ('20/01/2026', 'INV-1002', 'فاتورة مشتريات', 2500.5, None),
              ('27/01/2026', 'INV-1003', 'فاتورة مشتريات', 1500, None)]
    bal, built = 5000, []
    for d, ref, desc, dr, cr in chrono:
        bal += (dr or 0) - (cr or 0)
        built.append((d, ref, desc, dr, cr, bal))
    for r, row in enumerate(reversed(built), 6):
        for c, v in enumerate(row, 1):
            ws.cell(row=r, column=c, value=v)
    wb.save(os.path.join(HERE, 'fixture_supplier.xlsx'))


def csv_cp1256():
    text = 'التاريخ;رقم المستند;البيان;مدين;دائن;الرصيد\n;;رصيد افتتاحي;;100,00;100,00\n05/01/2026;F1;فاتورة مشتريات;;50,00;150,00\n06/01/2026;P1;سداد;20,00;;130,00\n'
    with open(os.path.join(HERE, 'fixture_cp1256.csv'), 'wb') as f:
        f.write(text.encode('cp1256'))


if __name__ == '__main__':
    ours(); supplier(); csv_cp1256()
    print('fixtures written')
